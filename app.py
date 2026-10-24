"""NSE Option Chain Dashboard — Streamlit app.

Uses nse_core (shared backend) for data and charts (Plotly builders) for
interactive visualizations.  All NSE scraping, parsing, formatting, analytics
and SQLite storage live in nse_core; all figure construction lives in charts.
This file owns only the Streamlit-specific UI layer.
"""

import datetime as dt
import io
import time
from pathlib import Path

import pandas as pd
import streamlit as st

import nse_core as nc
import charts

# ──────────────────────────────── page config ─────────────────────────────────
st.set_page_config(page_title="NSE Option Chain", layout="wide")
st.title("📋 NSE Option Chain Dashboard")

# ────────────────────── UI-specific constants & CSS ───────────────────────────

CSS = (
    "<style>"
    ".nse-wrap{max-height:640px;overflow:auto;border:1px solid #cfcfcf;background:#fff}"
    "table.nse{border-collapse:collapse;width:100%;font-family:Arial,Helvetica,sans-serif;"
    "font-size:13px;color:#222;background:#fff}"
    ".nse th{background:#3a2f86;color:#fff;font-weight:600;padding:8px 8px;"
    "border:1px solid #5b50a8;position:sticky;z-index:2;text-align:center;white-space:nowrap}"
    ".nse tr.g th{top:0;height:30px;box-sizing:border-box}"
    ".nse tr.h th{top:30px}"
    ".nse td{padding:8px 9px;text-align:center;border-bottom:1px solid #ececec;white-space:nowrap}"
    ".nse td.itm{background:#efe9cf}"
    ".nse td.k{font-weight:700;color:#1a4fa0;text-decoration:underline}"
    ".nse td.atm{background:#ffe066}"
    ".nse td.neg{color:#d32f2f}.nse td.pos{color:#1b873f}"
    "</style>"
)


# ────────────────────────── cached resource helpers ───────────────────────────

@st.cache_resource(show_spinner="Connecting to NSE …")
def _nse_client():
    return nc.NSE()


@st.cache_data(ttl=3600, show_spinner="Fetching expiry dates …")
def _expiries(symbol: str):
    return nc.get_expiries(_nse_client(), symbol)


# ─────────────────────── UI-specific display helpers ─────────────────────────

def _cell(v, name, dec, itm):
    cls = []
    if itm:
        cls.append("itm")
    if name == "CHNG":
        try:
            x = float(v)
        except (TypeError, ValueError):
            x = 0.0
        if x < 0:
            cls.append("neg")
        elif x > 0:
            cls.append("pos")
    return f'<td class="{" ".join(cls)}">{nc.fmt_in(v, dec)}</td>'


def _specs(fields):
    ce = [(n, d) for n, d in nc.CE_SPEC if n in fields]
    pe = [(n, d) for n, d in nc.PE_SPEC if n in fields]
    return ce, pe


def nse_html(df: pd.DataFrame, spot, fields, hl_atm=False) -> str:
    ce, pe = _specs(fields)
    atm = None
    if hl_atm and spot is not None and len(df):
        atm = df.loc[(df["STRIKE"] - spot).abs().idxmin(), "STRIKE"]
    head1 = (f'<tr class="g"><th colspan="{len(ce)}">CALLS</th><th></th>'
             f'<th colspan="{len(pe)}">PUTS</th></tr>')
    heads = ([n.replace("_", " ") for n, _ in ce] + ["STRIKE"]
             + [n.replace("_", " ") for n, _ in pe])
    head2 = '<tr class="h">' + "".join(f"<th>{h}</th>" for h in heads) + "</tr>"
    rows = []
    for _, r in df.iterrows():
        k = r["STRIKE"]
        itm_ce = spot is not None and k < spot
        itm_pe = spot is not None and k > spot
        tds = [_cell(r["CE_" + n], n, d, itm_ce) for n, d in ce]
        tds.append(f'<td class="{"k atm" if k == atm else "k"}">{nc.fmt_in(k, 2)}</td>')
        tds += [_cell(r["PE_" + n], n, d, itm_pe) for n, d in pe]
        rows.append("<tr>" + "".join(tds) + "</tr>")
    return (CSS + '<div class="nse-wrap"><table class="nse">' + head1 + head2
            + "".join(rows) + "</table></div>")


def nse_excel(df: pd.DataFrame, spot, fields, as_text: bool) -> bytes:
    """Excel in the NSE layout. as_text=True -> exactly what the website shows
    (Indian commas, '-' for zeros). as_text=False -> real numbers for calculations."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ce, pe = _specs(fields)
    n_ce, npe = len(ce), len(pe)
    total = n_ce + 1 + npe
    wb = Workbook()
    ws = wb.active
    ws.title = "OptionChain"
    hdr_fill = PatternFill("solid", fgColor="3A2F86")
    itm_fill = PatternFill("solid", fgColor="EFE9CF")
    white = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center")

    ws.cell(1, 1, "CALLS")
    if n_ce > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_ce)
    ws.cell(1, n_ce + 2, "PUTS")
    if npe > 1:
        ws.merge_cells(start_row=1, start_column=n_ce + 2, end_row=1, end_column=total)
    heads = ([n.replace("_", " ") for n, _ in ce] + ["STRIKE"]
             + [n.replace("_", " ") for n, _ in pe])
    for c, h in enumerate(heads, 1):
        ws.cell(2, c, h)
    for row in (1, 2):
        for c in range(1, total + 1):
            cell = ws.cell(row, c)
            cell.fill, cell.font, cell.alignment = hdr_fill, white, center

    specs = [("CE_", n, d) for n, d in ce] + [("", "STRIKE", 2)] + \
            [("PE_", n, d) for n, d in pe]
    for i, (_, r) in enumerate(df.iterrows(), start=3):
        k = r["STRIKE"]
        for c, (pre, n, d) in enumerate(specs, 1):
            v = r[pre + n] if pre else r[n]
            if as_text:
                cell = ws.cell(i, c, nc.fmt_in(v, d))
            else:
                try:
                    x = float(v)
                    x = None if x != x else x
                except (TypeError, ValueError):
                    x = None
                cell = ws.cell(i, c, x)
                cell.number_format = "#,##0.00" if d else "#,##0"
            cell.alignment = center
            if spot is not None and ((pre == "CE_" and k < spot) or (pre == "PE_" and k > spot)):
                cell.fill = itm_fill
    for c in range(1, total + 1):
        ws.column_dimensions[get_column_letter(c)].width = 13
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_excel(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="OptionChain")
    return buf.getvalue()


HISTORY_DIR = nc.APP_DIR / "history"


def log_snapshot(df, symbol, expiry, spot, ts):
    """Append one timestamped snapshot (all columns) to history/<symbol>_<expiry>.csv."""
    HISTORY_DIR.mkdir(exist_ok=True)
    path = HISTORY_DIR / f"{symbol}_{expiry}.csv"
    snap = df.copy()
    snap.insert(0, "nse_timestamp", ts)
    snap.insert(1, "spot", spot)
    snap.insert(0, "logged_at", dt.datetime.now().isoformat(timespec="seconds"))
    snap.to_csv(path, mode="a", header=not path.exists(), index=False)
    return path


# ─────────────────────────────── sidebar ──────────────────────────────────────
cfg = nc.load_config()

with st.sidebar:
    st.header("⚙ Settings")
    kinds = ["Indices", "Equity"]
    kind = st.radio("Type", kinds, index=kinds.index(cfg["kind"]) if cfg["kind"] in kinds else 0,
                    horizontal=True)
    if kind == "Indices":
        sym0 = cfg["symbol"] if cfg["symbol"] in nc.INDICES else "NIFTY"
        symbol = st.selectbox("Symbol", nc.INDICES, index=nc.INDICES.index(sym0))
    else:
        symbol = st.text_input(
            "Stock symbol", cfg["symbol"] if cfg["kind"] == "Equity" else "RELIANCE"
        ).strip().upper()

    try:
        expiries = _expiries(symbol)
    except Exception as e:  # noqa: BLE001
        expiries = []
        st.warning(f"Could not load expiries: {e}")
    expiry = (
        st.selectbox("Expiry", expiries)
        if expiries
        else st.text_input("Expiry (e.g. 27-Oct-2026)")
    )

    modes = ["NSE website format", "Custom columns"]
    view_mode = st.radio("Table format", modes,
                         index=modes.index(cfg["view_mode"]) if cfg["view_mode"] in modes else 0)
    fields, cols, hl_atm = nc.ALL_NSE_FIELDS, cfg["custom_cols"], False
    if view_mode == modes[1]:
        cols = st.multiselect(
            "Columns", list(nc.FIELDS.values()),
            default=[c for c in cfg["custom_cols"] if c in nc.FIELDS.values()] or nc.DEFAULT_COLS,
        ) or nc.DEFAULT_COLS
    else:
        fields = st.multiselect(
            "Fields to show (same on CALLS and PUTS)", nc.ALL_NSE_FIELDS,
            default=[f for f in cfg["fields"] if f in nc.ALL_NSE_FIELDS] or nc.ALL_NSE_FIELDS,
            format_func=lambda s: s.replace("_", " "),
        ) or nc.ALL_NSE_FIELDS
        hl_atm = st.checkbox("Highlight ATM strike", value=bool(cfg["highlight_atm"]))
    opts = [0, 5, 10, 15, 20, 30]
    atm_n = st.select_slider("Strikes around ATM", opts,
                             value=cfg["atm_n"] if cfg["atm_n"] in opts else 0,
                             help="0 = all strikes (same as the NSE page)")
    st.divider()
    auto = st.checkbox("Auto-refresh", value=bool(cfg["auto"]))
    iv = min(300, max(15, int(cfg["interval"])))
    interval = st.slider("Every (seconds)", 15, 300, iv - iv % 15 or 15, step=15)
    log_hist = st.checkbox("Log snapshots to CSV & SQLite", value=bool(cfg["log"]))
    st.button("🔄 Refresh now")
    if st.button("💾 Save as default settings"):
        nc.save_config({
            "kind": kind, "symbol": symbol, "view_mode": view_mode,
            "fields": fields, "custom_cols": cols, "atm_n": atm_n,
            "highlight_atm": hl_atm, "auto": auto, "interval": interval, "log": log_hist,
        })
        st.success("Saved to config.json")

# ──────────────────────────── data fetching ───────────────────────────────────

if not expiry:
    st.info("Pick or type an expiry to continue.")
    st.stop()

try:
    with st.spinner("Fetching from NSE..."):
        raw = nc.get_chain(_nse_client(), symbol, expiry, kind)
    full, spot, ts = nc.parse(raw)
except Exception as e:  # noqa: BLE001
    st.error(f"❌ Fetch/parse failed: {e}")
    st.caption("Try Refresh now. If it keeps failing, NSE may have changed the API "
               "or blocked the request. Check the Network tab on the NSE page.")
    st.stop()

trimmed = nc.trim_atm(full, spot, atm_n)
summary = nc.summarize(full, spot)

# ─────────────────────────── main area: tabs ──────────────────────────────────
tab_chain, tab_charts, tab_history = st.tabs(
    ["📊 Option Chain", "📈 Charts", "📜 History"]
)

# ·································· Tab 1 ··································
with tab_chain:
    # metrics row
    m1, m2, m3, m4, m5 = st.columns(5)
    m1.metric(f"{symbol} spot", nc.fmt_in(spot, 2) if spot else "—")
    m2.metric("PCR (OI)", f"{summary['pcr']:.2f}" if summary["pcr"] == summary["pcr"] else "—")
    m3.metric("Total CE OI", nc.fmt_in(summary["ce_oi"]))
    m4.metric("Total PE OI", nc.fmt_in(summary["pe_oi"]))
    m5.metric("Max Pain", nc.fmt_in(summary["max_pain"], 2) if summary["max_pain"] else "—")

    mkt = "🟢 Market open" if nc.market_open() else "🔴 Market closed"
    st.caption(f"NSE data timestamp: {ts}   |   Expiry: {expiry}   |   {mkt}")

    # table
    if view_mode == "NSE website format":
        st.markdown(nse_html(trimmed, spot, fields, hl_atm), unsafe_allow_html=True)
    else:
        ce_view = ["CE_" + c for c in cols]
        pe_view = ["PE_" + c for c in reversed(cols)]
        view = trimmed[ce_view + ["STRIKE"] + pe_view]

        def hl(row):
            if spot is not None and abs(row["STRIKE"] - spot) == (view["STRIKE"] - spot).abs().min():
                return ["background-color: #fff3bf; color: #000"] * len(row)
            return [""] * len(row)

        st.dataframe(view.style.apply(hl, axis=1), use_container_width=True,
                     hide_index=True, height=560)

    # downloads
    c1, c2, c3 = st.columns(3)
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if view_mode == "NSE website format":
        c1.download_button("⬇ Download Excel (NSE format)",
                           nse_excel(trimmed, spot, fields, as_text=True),
                           file_name=f"{symbol}_{expiry}_nse_format.xlsx", mime=xlsx)
        c2.download_button("⬇ Download Excel (numbers)",
                           nse_excel(trimmed, spot, fields, as_text=False),
                           file_name=f"{symbol}_{expiry}_numbers.xlsx", mime=xlsx,
                           help="Real numeric cells, so you can sort and use formulas")
    else:
        c1.download_button("⬇ Download Excel",
                           to_excel(trimmed[ce_view + ["STRIKE"] + pe_view]),
                           file_name=f"{symbol}_{expiry}_option_chain.xlsx", mime=xlsx)

    # logging (CSV + SQLite)
    if log_hist:
        key = (symbol, expiry, ts)
        if st.session_state.get("last_logged") != key:
            p = log_snapshot(full, symbol, expiry, spot, ts)
            nc.save_snapshot(full, symbol, expiry, spot, ts)
            st.session_state["last_logged"] = key
            st.sidebar.success(f"Logged → {p.name} & SQLite")
        hp = HISTORY_DIR / f"{symbol}_{expiry}.csv"
        if hp.exists():
            c3.download_button("⬇ Download history CSV", hp.read_bytes(),
                               file_name=hp.name, mime="text/csv")

# ·································· Tab 2 ··································
with tab_charts:
    if trimmed.empty:
        st.info("No data to chart.")
    else:
        r1c1, r1c2 = st.columns(2)
        with r1c1:
            st.plotly_chart(charts.oi_by_strike(trimmed, spot), use_container_width=True)
        with r1c2:
            st.plotly_chart(charts.change_in_oi(trimmed, spot), use_container_width=True)

        r2c1, r2c2 = st.columns(2)
        with r2c1:
            st.plotly_chart(charts.volume_by_strike(trimmed, spot), use_container_width=True)
        with r2c2:
            st.plotly_chart(charts.iv_smile(trimmed, spot), use_container_width=True)

        st.plotly_chart(charts.premium_by_strike(trimmed, spot), use_container_width=True)

# ·································· Tab 3 ··································
with tab_history:
    days = nc.history_days(symbol, expiry)
    if not days:
        st.info(
            "No intraday history recorded yet for this symbol / expiry.\n\n"
            "To start collecting snapshots:\n"
            "- Run **collect.bat** in the background, **or**\n"
            "- Enable **Log snapshots** in the sidebar and turn on **Auto-refresh**."
        )
    else:
        sel_day = st.selectbox("Select day", days, index=0)
        h = nc.load_history(symbol, expiry, sel_day)

        if h.empty:
            st.warning("No rows for the selected day.")
        else:
            hc1, hc2 = st.columns(2)
            with hc1:
                st.plotly_chart(charts.spot_pcr(h), use_container_width=True)
            with hc2:
                st.plotly_chart(charts.total_oi(h), use_container_width=True)

            strikes = sorted(h["STRIKE"].dropna().unique())
            if strikes:
                sel_strike = st.selectbox("Select strike for detail",
                                          strikes, index=len(strikes) // 2)
                sc1, sc2 = st.columns(2)
                with sc1:
                    st.plotly_chart(charts.strike_premium(h, sel_strike),
                                   use_container_width=True)
                with sc2:
                    st.plotly_chart(charts.strike_oi(h, sel_strike),
                                   use_container_width=True)

# ──────────────────────────── auto-refresh ────────────────────────────────────
if auto:
    time.sleep(interval)
    st.rerun()
